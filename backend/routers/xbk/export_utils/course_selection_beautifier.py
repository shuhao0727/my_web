"""
学生选课表美化器

功能：专门负责美化已导出的学生选课表Excel文件
设计原则：
1. 单一职责：仅负责Excel文件的美化（样式、验证、保护等）
2. 无GUI依赖：移除PyQt6依赖，纯后端美化
3. 与数据导出器解耦：接收文件路径，处理后保存
"""
import openpyxl
from openpyxl.styles import Border, Protection, Alignment, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import logging
import os

# 设置日志记录
logger = logging.getLogger(__name__)

def auto_adjust_column_width(sheet, special=False):
    """根据内容自动调整列宽"""
    for col in sheet.iter_cols():
        max_length = 0
        column = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception as e:
                    logger.warning(f"调整列宽时出错 {cell.coordinate}: {e}")
        if special:
            adjusted_width = (max_length + 5) * 1.5
        else:
            adjusted_width = (max_length + 5) * 1.2
        sheet.column_dimensions[column].width = adjusted_width

def auto_adjust_row_height(sheet):
    """根据内容自动调整行高"""
    for row in sheet.rows:
        max_height = 15
        for cell in row:
            if cell.value:
                line_count = str(cell.value).count('\n') + 1
                content_length = len(str(cell.value))
                if content_length > 50:
                    line_count += (content_length - 50) // 50
                max_height = max(max_height, line_count * 15)
        sheet.row_dimensions[row[0].row].height = max_height

def convert_to_number(sheet, column_header):
    """将指定列的数据转换为数字类型"""
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if sheet.cell(row=1, column=cell.column).value == column_header:
                try:
                    cell.value = int(cell.value)
                except (ValueError, TypeError):
                    cell.value = None

def convert_course_code_to_number(sheet):
    """将课程代码列转换为数字类型"""
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if sheet.cell(row=1, column=cell.column).value == "课程代码":
                try:
                    cell.value = int(cell.value)
                except (ValueError, TypeError):
                    cell.value = None

def get_course_code_limits(course_catalog_sheet):
    """获取课程代码和限报人数"""
    course_codes = []
    limits = {}
    for row in course_catalog_sheet.iter_rows(min_row=2):
        course_code = None
        limit = None
        for cell in row:
            if course_catalog_sheet.cell(row=1, column=cell.column).value == "课程代码":
                course_code = cell.value
            if course_catalog_sheet.cell(row=1, column=cell.column).value == "各班限报人数":
                limit = cell.value
        if course_code is not None and limit is not None:
            course_codes.append(str(course_code))
            limits[str(course_code)] = int(limit)
    return course_codes, limits

def set_data_validation(sheet, column_letter, valid_course_codes, catalog_sheet_name):
    """设置数据验证"""
    for row in range(2, sheet.max_row + 1):
        if valid_course_codes:
            # 将课程代码转换为整数并获取最小值和最大值
            int_codes = []
            for code in valid_course_codes:
                try:
                    int_codes.append(int(code))
                except (ValueError, TypeError):
                    continue
            
            if int_codes:
                min_code = min(int_codes)
                max_code = max(int_codes)
                formula = f'=AND(ISNUMBER({column_letter}{row}), {column_letter}{row}>={min_code}, {column_letter}{row}<={max_code}, COUNTIF({column_letter}:{column_letter}, {column_letter}{row})<=VLOOKUP({column_letter}{row}, {catalog_sheet_name}!A:E, 4, 0))'
            else:
                formula = f'=ISNUMBER({column_letter}{row})'
        else:
            formula = f'=ISNUMBER({column_letter}{row})'
        
        dv = DataValidation(type="custom", formula1=formula, showErrorMessage=True, 
                           errorTitle="错误", error="只能输入有效的课程代码，并且数量不能超过限制")
        sheet.add_data_validation(dv)
        dv.add(f'{column_letter}{row}')

def lock_worksheet(sheet, protect_all=True, unlock_column=None):
    """锁定工作表"""
    sheet.protection.sheet = True
    # 注意：实际使用时可以设置密码，这里暂时使用空密码
    sheet.protection.set_password('')
    
    if protect_all:
        for row in sheet.iter_rows():
            for cell in row:
                cell.protection = Protection(locked=True)
    
    if unlock_column:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if sheet.cell(row=1, column=cell.column).value == unlock_column:
                    cell.protection = Protection(locked=False)

def apply_excel_style(sheet, header_fill_color="CCE5FF"):
    """应用Excel样式"""
    # 设置表头样式
    header_fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
    header_font = Font(bold=True, size=11, color="000000")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin'))
    
    # 应用表头样式
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # 设置数据单元格样式
    data_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    alternate_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    data_font = Font(size=10)
    
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if row_idx % 2 == 0:
                cell.fill = alternate_fill
            else:
                cell.fill = data_fill
    
    # 自动调整列宽和行高
    auto_adjust_column_width(sheet)
    auto_adjust_row_height(sheet)

def beautify_course_selection_file(file_path: str):
    """
    美化学生选课表Excel文件
    
    参数:
        file_path: Excel文件路径
        
    返回:
        bool: 美化是否成功
    """
    try:
        if not file_path.endswith('.xlsx'):
            logger.error(f"仅支持.xlsx文件: {file_path}")
            return False
        
        workbook = openpyxl.load_workbook(file_path)
        
        # 获取校本课程目录中的课程代码和各班限报人数
        if "校本课程目录" not in workbook.sheetnames:
            logger.error("文件缺少'校本课程目录'工作表")
            return False
            
        course_catalog_sheet = workbook["校本课程目录"]
        valid_course_codes, course_code_limits = get_course_code_limits(course_catalog_sheet)
        
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            
            # 将"各班限报人数"列数据修改为数字
            convert_to_number(ws, "各班限报人数")
            
            # 将"课程代码"列数据修改为数字（校本课程目录表格）
            if "校本课程目录" in sheet_name:
                convert_course_code_to_number(ws)
                # 添加表头
                ws.insert_rows(1)
                # 根据列数调整合并单元格范围（删除"对选课学生的要求"列后，现在是5列）
                col_count = ws.max_column
                if col_count >= 5:
                    ws.merge_cells(f'A1:E1')
                elif col_count > 0:
                    ws.merge_cells(f'A1:{get_column_letter(col_count)}1')
                
                ws['A1'] = "江苏省昆山中学校本课程目录"
                ws['A1'].font = Font(size=14, bold=True)
                ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
            
            # 针对校本课程目录表格使用特殊宽度调整
            if "校本课程目录" in sheet_name:
                auto_adjust_column_width(ws, special=True)
                # 锁定整个表格
                lock_worksheet(ws, protect_all=True)
                # 应用样式
                apply_excel_style(ws, header_fill_color="CCE5FF")
            else:
                auto_adjust_column_width(ws)
                # 设置"课程代码"列的有效性
                for col in ws.iter_cols(min_row=1, max_row=1):
                    for cell in col:
                        if cell.value == "课程代码":
                            column_letter = get_column_letter(cell.column)
                            set_data_validation(ws, column_letter, valid_course_codes, "校本课程目录")
                # 锁定除"课程代码"列外的所有列
                lock_worksheet(ws, protect_all=True, unlock_column="课程代码")
                # 应用样式
                apply_excel_style(ws, header_fill_color="CCE5FF")
        
        workbook.save(file_path)
        logger.info(f"学生选课表美化成功: {file_path}")
        return True
        
    except Exception as e:
        logger.error(f"美化学生选课表失败: {str(e)}")
        return False

def test_beautify():
    """测试美化功能"""
    try:
        # 先导出测试文件
        from .course_selection_exporter import export_course_selection
        test_file = export_course_selection(
            year=2025,
            grade="高一",
            year_start=2024,
            year_end=2025
        )
        
        if test_file and os.path.exists(test_file):
            print(f"测试文件: {test_file}")
            print(f"美化前文件大小: {os.path.getsize(test_file)} 字节")
            
            # 应用美化
            result = beautify_course_selection_file(test_file)
            
            if result:
                print("美化成功!")
                print(f"美化后文件大小: {os.path.getsize(test_file)} 字节")
                
                # 验证美化效果
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(test_file)
                    ws = wb["校本课程目录"]
                    
                    # 检查是否有表头
                    if ws['A1'].value == "江苏省昆山中学校本课程目录":
                        print("✓ 表头添加成功")
                    
                    # 检查样式
                    if ws['A2'].fill.start_color.rgb == "CCE5FF":
                        print("✓ 表头样式应用成功")
                    
                    # 检查数据验证
                    if len(ws.data_validations.dataValidation) > 0:
                        print("✓ 数据验证设置成功")
                    
                    wb.close()
                except Exception as e:
                    print(f"验证美化效果时出错: {e}")
                
                return True
            else:
                print("美化失败")
                return False
        else:
            print("测试文件创建失败")
            return False
            
    except Exception as e:
        print(f"测试美化功能失败: {e}")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    # 运行测试
    print("=== 测试学生选课表美化功能 ===")
    test_beautify()
